"""Write one Excel workbook listing every KV extractor key and its box expansion.

Sheets (one per extractor):
  DOB | Member_ID | Member_Name | Provider_Name | Electronic_Signature

Columns:
  Key | Box_Expansion | Suggested_Expansion | Reason_For_Suggested

Usage (from repo root):
  .\\.venv\\Scripts\\python.exe write_key_expansions.py
"""

from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

ROOT = Path(__file__).resolve().parent
KV = ROOT / "KV_Extraction"
OUT = ROOT / "Key_Box_Expansions.xlsx"

# Same rule as KV_Extraction/Util/window.py expand_for_key
SHORT_EXPAND = 6.0
LONG_EXPAND = 5.0

SHEETS = [
    ("DOB", KV / "Member_DOB" / "keys.json"),
    ("Member_ID", KV / "Member_ID" / "keys.json"),
    ("Member_Name", KV / "Member_Name" / "keys.json"),
    ("Provider_Name", KV / "Provider_Name" / "keys.json"),
    ("Electronic_Signature", KV / "Electronic_Signature" / "keys.json"),
]

HEADERS = ["Key", "Box_Expansion", "Suggested_Expansion", "Reason_For_Suggested"]


def expand_for_key(key: str) -> str:
    scale = SHORT_EXPAND if len(key.strip()) <= 5 else LONG_EXPAND
    return f"{scale:g}x"


def load_keys(path: Path) -> list[str]:
    data = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(data, list):
        raise SystemExit(f"Expected a JSON list in {path}")
    return [str(item).strip() for item in data if str(item).strip()]


def main() -> int:
    book = Workbook()
    book.remove(book.active)

    header_font = Font(name="Arial", bold=True, size=11)
    cell_font = Font(name="Arial", size=11)
    for title, keys_path in SHEETS:
        if not keys_path.is_file():
            raise SystemExit(f"Missing keys file: {keys_path}")
        sheet = book.create_sheet(title)
        for col, header in enumerate(HEADERS, start=1):
            cell = sheet.cell(1, col, header)
            cell.font = header_font
            cell.alignment = Alignment(horizontal="left")
        for row, key in enumerate(load_keys(keys_path), start=2):
            sheet.cell(row, 1, key).font = cell_font
            sheet.cell(row, 2, expand_for_key(key)).font = cell_font
            sheet.cell(row, 3, "").font = cell_font
            sheet.cell(row, 4, "").font = cell_font
        sheet.column_dimensions["A"].width = 28
        sheet.column_dimensions["B"].width = 16
        sheet.column_dimensions["C"].width = 20
        sheet.column_dimensions["D"].width = 36
        sheet.freeze_panes = "A2"

    book.save(OUT)
    print(OUT)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
