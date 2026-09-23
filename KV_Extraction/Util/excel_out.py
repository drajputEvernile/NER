"""Build the combined Extraction + Extraction Summary Excel workbook."""

from __future__ import annotations

import json
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils.dataframe import dataframe_to_rows

FIELD_SPECS = [
    {
        "prefix": "dob",
        "detail": "extraction_dob.csv",
        "value_col": "Value",
        "summary_col": "DOB",
        "extra_value_cols": (),
    },
    {
        "prefix": "ID",
        "detail": "extraction_id.csv",
        "value_col": "Value",
        "summary_col": "ID",
        "extra_value_cols": (),
    },
    {
        "prefix": "MName",
        "detail": "extraction_name.csv",
        "value_col": "Value",
        "summary_col": "MName",
        "extra_value_cols": (),
    },
    {
        "prefix": "PName",
        "detail": "extraction_provider.csv",
        "value_col": "Value",
        "summary_col": "PName",
        "extra_value_cols": (),
    },
    {
        "prefix": "ESig",
        "detail": "extraction_esig.csv",
        "value_col": "ProviderName",
        "summary_col": "ESign_Provider",
        "extra_value_cols": (("SignatureDate", "SignatureDate"),),
        "summary_extra": (("ESign_Date", "SignatureDate"),),
    },
]

HIT_ATTRS = [
    ("Key", "Key", False),
    ("Region", "Region", True),
    ("Sentence", "Sentence", True),
    ("Ner_Text", "Ner_Text", True),
    ("Value", "Value", True),
    ("Score", "Score", True),
    ("Accepted", "Accepted", True),
    ("Selected", "Selected", True),
    ("Source", "Source", True),
    ("Accuracy", "Accuracy", True),
]


def _list_cell(items: list[str]) -> str:
    return json.dumps(items, ensure_ascii=False)


def _yes_no(value: Any) -> str:
    text = str(value or "").strip().casefold()
    if text in {"yes", "true", "1"}:
        return "yes"
    if text in {"no", "false", "0"}:
        return "no"
    return str(value or "")


def _load_detail(path: Path) -> pd.DataFrame:
    if not path.is_file():
        return pd.DataFrame()
    frame = pd.read_csv(path, dtype=str).fillna("")
    return frame


def _page_key(row: pd.Series) -> tuple[str, str, str]:
    return (
        str(row.get("RecordId") or ""),
        str(row.get("FileName") or ""),
        str(row.get("PageNumber") or ""),
    )


def _field_columns(prefix: str, *, esig: bool = False) -> list[str]:
    cols = [
        f"{prefix}_Key",
        f"{prefix}_Region",
        f"{prefix}_Sentence",
        f"{prefix}_Ner_Text",
    ]
    if esig:
        cols.extend([f"{prefix}_ProviderName", f"{prefix}_SignatureDate"])
    else:
        cols.append(f"{prefix}_Value")
    cols.extend(
        [
            f"{prefix}_Score",
            f"{prefix}_Accepted",
            f"{prefix}_Selected",
            f"{prefix}_Source",
            f"{prefix}_Accuracy",
        ]
    )
    return cols


def _pack_hits(group: pd.DataFrame, *, value_col: str, esig: bool = False) -> dict[str, str]:
    prefix = "ESig" if esig else None
    # Caller passes prefix via columns built outside; here we only pack values.
    keys = [str(row.get("Key") or "") for _, row in group.iterrows()]
    out: dict[str, list[str]] = {
        "Key": keys,
        "Region": [],
        "Sentence": [],
        "Ner_Text": [],
        "Value": [],
        "ProviderName": [],
        "SignatureDate": [],
        "Score": [],
        "Accepted": [],
        "Selected": [],
        "Source": [],
        "Accuracy": [],
    }
    for _, row in group.iterrows():
        key = str(row.get("Key") or "")
        out["Region"].append(f"{key}: {row.get('Region') or ''}")
        out["Sentence"].append(f"{key}: {row.get('Sentence') or ''}")
        out["Ner_Text"].append(f"{key}: {row.get('Ner_Text') or ''}")
        if esig:
            out["ProviderName"].append(f"{key}: {row.get('ProviderName') or ''}")
            out["SignatureDate"].append(f"{key}: {row.get('SignatureDate') or ''}")
        else:
            out["Value"].append(f"{key}: {row.get(value_col) or ''}")
        out["Score"].append(f"{key}: {row.get('Score') or ''}")
        out["Accepted"].append(f"{key}: {_yes_no(row.get('Accepted'))}")
        out["Selected"].append(f"{key}: {_yes_no(row.get('Selected'))}")
        out["Source"].append(f"{key}: {row.get('Source') or ''}")
        out["Accuracy"].append(f"{key}: ")
    del prefix
    return {name: _list_cell(values) for name, values in out.items()}


def _selected_summary(group: pd.DataFrame, value_col: str) -> str:
    if group.empty:
        return ""
    selected = group[group["Selected"].astype(str).str.casefold().isin(["yes", "true", "1"])]
    pool = selected if not selected.empty else group
    accepted = pool[pool["Accepted"].astype(str).str.casefold().isin(["yes", "true", "1"])]
    pool = accepted if not accepted.empty else pool
    values = [str(v).strip() for v in pool[value_col].tolist() if str(v).strip()]
    if not values:
        return ""
    # Majority then first.
    counts: dict[str, int] = {}
    for value in values:
        counts[value.casefold()] = counts.get(value.casefold(), 0) + 1
    best = max(counts.values())
    for value in values:
        if counts[value.casefold()] == best:
            return value
    return values[0]


def build_extraction_frames(
    detail_dirs: dict[str, Path],
    page_counts: dict[str, int],
    time_seconds: dict[str, float],
) -> tuple[pd.DataFrame, pd.DataFrame]:
    """Assemble page-level Extraction + document-level Summary frames."""
    loaded: dict[str, pd.DataFrame] = {}
    pages: set[tuple[str, str, str]] = set()
    for spec in FIELD_SPECS:
        folder = detail_dirs.get(spec["prefix"])
        path = (folder / spec["detail"]) if folder else None
        frame = _load_detail(path) if path else pd.DataFrame()
        loaded[spec["prefix"]] = frame
        if not frame.empty:
            for _, row in frame.iterrows():
                pages.add(_page_key(row))

    extraction_rows: list[dict[str, Any]] = []
    for record_id, file_name, page_number in sorted(pages, key=lambda item: (item[0], int(item[2] or 0), item[1])):
        row: dict[str, Any] = {
            "RecordId": record_id,
            "PageCount": str(page_counts.get(record_id, "")),
            "FileName": file_name,
            "PageNumber": page_number,
        }
        for spec in FIELD_SPECS:
            prefix = spec["prefix"]
            frame = loaded[prefix]
            if frame.empty:
                group = frame
            else:
                mask = (
                    (frame["RecordId"].astype(str) == record_id)
                    & (frame["FileName"].astype(str) == file_name)
                    & (frame["PageNumber"].astype(str) == page_number)
                )
                group = frame.loc[mask]
            esig = prefix == "ESig"
            packed = _pack_hits(group, value_col=spec["value_col"], esig=esig) if len(group) else None
            if packed is None:
                for col in _field_columns(prefix, esig=esig):
                    row[col] = ""
                continue
            row[f"{prefix}_Key"] = packed["Key"]
            row[f"{prefix}_Region"] = packed["Region"]
            row[f"{prefix}_Sentence"] = packed["Sentence"]
            row[f"{prefix}_Ner_Text"] = packed["Ner_Text"]
            if esig:
                row[f"{prefix}_ProviderName"] = packed["ProviderName"]
                row[f"{prefix}_SignatureDate"] = packed["SignatureDate"]
            else:
                row[f"{prefix}_Value"] = packed["Value"]
            row[f"{prefix}_Score"] = packed["Score"]
            row[f"{prefix}_Accepted"] = packed["Accepted"]
            row[f"{prefix}_Selected"] = packed["Selected"]
            row[f"{prefix}_Source"] = packed["Source"]
            row[f"{prefix}_Accuracy"] = packed["Accuracy"]
        extraction_rows.append(row)

    extraction_cols = ["RecordId", "PageCount", "FileName", "PageNumber"]
    for spec in FIELD_SPECS:
        extraction_cols.extend(_field_columns(spec["prefix"], esig=spec["prefix"] == "ESig"))
    extraction = pd.DataFrame(extraction_rows, columns=extraction_cols)

    summary_rows: list[dict[str, Any]] = []
    for record_id in sorted(page_counts):
        summary: dict[str, Any] = {
            "RecordId": record_id,
            "PageCount": str(page_counts.get(record_id, "")),
            "DOB": "",
            "ID": "",
            "MName": "",
            "PName": "",
            "ESign_Provider": "",
            "ESign_Date": "",
            "TimeSeconds": (
                f"{time_seconds[record_id]:.3f}" if record_id in time_seconds else ""
            ),
        }
        for spec in FIELD_SPECS:
            frame = loaded[spec["prefix"]]
            if frame.empty:
                continue
            group = frame[frame["RecordId"].astype(str) == record_id]
            summary[spec["summary_col"]] = _selected_summary(group, spec["value_col"])
            for summary_name, source_col in spec.get("summary_extra") or ():
                summary[summary_name] = _selected_summary(group, source_col)
        summary_rows.append(summary)

    summary = pd.DataFrame(
        summary_rows,
        columns=[
            "RecordId",
            "PageCount",
            "DOB",
            "ID",
            "MName",
            "PName",
            "ESign_Provider",
            "ESign_Date",
            "TimeSeconds",
        ],
    )
    return extraction, summary


def write_workbook(path: Path, extraction: pd.DataFrame, summary: pd.DataFrame) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    book = Workbook()
    book.remove(book.active)
    header_font = Font(name="Arial", bold=True, size=11)
    cell_font = Font(name="Arial", size=10)
    for title, frame in (("Extraction", extraction), ("Extraction Summary", summary)):
        sheet = book.create_sheet(title)
        for row_index, row in enumerate(dataframe_to_rows(frame, index=False, header=True), start=1):
            for col_index, value in enumerate(row, start=1):
                cell = sheet.cell(row_index, col_index, value)
                cell.font = header_font if row_index == 1 else cell_font
        sheet.freeze_panes = "A2"
    book.save(path)
    return path
